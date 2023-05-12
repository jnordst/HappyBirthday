# Course Project
# [COMP 1112] Document Automation Python
# Jacob Nordstrom
# 200368110
# 03/13/2023
# Search through sheet of clients and get their birthday information
# If application finds a client that matches sends them a personalized email

#-------------------------------------
# modules required for sending emails
import smtplib 
from email.message import EmailMessage
import imghdr
#-------------------------------------
import openpyxl
import datetime
import re
import os

# securely hides address and password in enivornment variables on computer rather than in python file
EMAIL_ADDRESS = os.environ.get('EMAIL_USER')
EMAIL_PASSWORD = os.environ.get('EMAIL_PASS')

# opens patients.xlsx and sets an active sheet
workbook = openpyxl.load_workbook("patients.xlsx")
sheet = workbook.active
activeRow = ''

# user must input their password to run the program
# iterates through a while loop until they provide the correct password
def StartApplication():
    validInput = False
    while not(validInput):
        userInput = input("Please enter your password")
        if (userInput == EMAIL_PASSWORD):
            validInput = True
        else:
            print("Password Incorrect")
            continue
    
    CheckForBirthdays()
    print("Complete")

# iterates through the worksheet
# gets the birthday information of each row and sends it to be verified by isBirthday()
# if it's the clients birthday sends the email address to be verified by isValidEmail()
# if email is valid, sends an email to the patient
def CheckForBirthdays():
    for row in range(1, sheet.max_row): # start at row 1 and go to sheets max rows
        birthDate = sheet['C' + str(row + 1)].value # value of sheet[C-row] is a mm-dd-yyyy formated birth date
        # check if it's the clients birthday
        if (isBirthday(birthDate)):
            global activeRow
            activeRow = str(row + 1)
            emailAddress = sheet['E' + activeRow].value

            # check if the email address is valid
            if (isValidEmail(emailAddress)):
                print(f"Sending email to {sheet['A' + activeRow].value} {sheet['B' + activeRow].value}")
                sendEmail(emailAddress)

# validates that the email address in the spreadsheet is valid
# if the email address is empty, notify user and return False but do not throw an exception since not every patient supplies an email address
# if the email address is valid and not None, return True
# if the email address is not None but not valid, raise an exception to notify the user that the email address is invalid
def isValidEmail(emailAddress):
    if (emailAddress is None):
        print(f"Email at E{activeRow} is empty")
        return False
    
    pattern = r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$' # generic email validation from https://www.c-sharpcorner.com/article/how-to-validate-an-email-address-in-python/
    match = re.match(pattern, emailAddress) is not None

    if (match):
        return True
    
    raise Exception(f"Email at E{activeRow} is not a valid email address")

# returns a boolean indicating whether the patients birthday matches todays date
# argument should be a string formatted as mm-dd-yyyy
def isBirthday(birthday):
    today = datetime.datetime.now()
    format = today.strftime("%m-%d")
    pattern = (format)
    regex = re.compile(pattern)
    match = regex.search(birthday)
    if (match is not None):
        return True
    return False

# open the body text file
# slices placeholders inside the body text file with personalized data from the patients spreadsheet
# return fileContent after personlizing it for the patient
def getBody():
    # using with open() opens a file and will automatically close it
    # body.txt is a textfile containing placeholder text, located in the same directory as the .py file
    with open("body.txt", "r") as file:
        fileContent = file.read()

    # create a dictionary to hold the regex pattern and replacement
    # the key is the pattern regex will search for
    # the value is the replacement
    regex = {
        "FIRST_NAME": sheet[f"A{activeRow}"].value,
        "LAST_NAME": sheet[f"B{activeRow}"].value,
        "LAST_VISIT": sheet[f"K{activeRow}"].value,
    }

    # iterate through the dictionary and using re.sub() to replace any patterns with personalized information from the spreadsheet
    for key, value in regex.items(): # specify the key, value and regex.items() to access both the key and value of the dictionary
        if value is not None:
            pattern = re.compile(key)
            fileContent = re.sub(pattern, str(value), fileContent) # re.sub() slices the string and inserts the new string (value). cast as str incase value isn't a string
        else:
            fileContent = re.sub(key, "", fileContent) # in case the spreadsheet data is empty, replace the value with and empty string

    # return the edited content
    return fileContent

# using email messsage and smptlib module to set up and send an email
# credits to Corey Schafer for explaining how to send emails through python
# https://www.youtube.com/watch?v=JRCJ6RtE3xU
def sendEmail(emailAddress):
    msg = EmailMessage() # creates a new EmailMessage object
    msg['Subject'] = "Happy Birthday!"
    msg['From'] = EMAIL_ADDRESS
    msg['To'] = emailAddress
    msg.set_content(getBody())

    # open the image file
    # imghdr module gets the file type (.png, .jpg etc)
    with open("birthday-card.png", 'rb') as file:
        file_data = file.read()
        file_type = imghdr.what(file.name)
        file_name = file.name

    # add the image as an attachment
    msg.add_attachment(file_data, maintype = "image", subtype = file_type, filename = file_name)

    with smtplib.SMTP_SSL('smtp.gmail.com', 465) as smtp: # estabolishes a secure smtp connection via port 465
        smtp.login(EMAIL_ADDRESS, EMAIL_PASSWORD) # login to email address
        smtp.send_message(msg) # send the email
   
StartApplication()